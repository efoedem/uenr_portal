"""Custom staff dashboard: students, dues publishing, payments and reports."""
from decimal import Decimal
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import user_passes_test
from django.contrib.auth.models import User
from django.db.models import Count, Q, Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from .models import LEVELS, PROGRAMMES, Due, DueTemplate, Payment, Student

staff_required = user_passes_test(lambda u: u.is_active and u.is_staff, login_url="/login/")

PROGRAMME_KEYS = dict(PROGRAMMES)
LEVEL_KEYS = dict(LEVELS)
DEFAULT_PASSWORD = "uenr@1234"


def _programme_from_text(text):
    t = (text or "").strip().lower()
    if not t:
        return ""
    if t in PROGRAMME_KEYS:
        return t
    if "comp" in t or t == "coe":
        return "computer_engineering"
    if "elec" in t or t == "eee":
        return "electrical_electronics"
    return ""


def _level_from_text(text):
    t = str(text or "").strip().replace("level", "").replace("Level", "").strip()
    t = t.split(".")[0]
    return t if t in LEVEL_KEYS else ""


def _filtered_students(request):
    qs = Student.objects.select_related("user")
    level = request.GET.get("level", "")
    programme = request.GET.get("programme", "")
    q = request.GET.get("q", "").strip()
    if level in LEVEL_KEYS:
        qs = qs.filter(level=level)
    if programme in PROGRAMME_KEYS:
        qs = qs.filter(programme=programme)
    if q:
        qs = qs.filter(
            Q(student_id__icontains=q)
            | Q(user__first_name__icontains=q)
            | Q(user__last_name__icontains=q)
        )
    return qs


def _filtered_payments(request):
    qs = Payment.objects.select_related("student", "student__user")
    level = request.GET.get("level", "")
    programme = request.GET.get("programme", "")
    method = request.GET.get("method", "")
    q = request.GET.get("q", "").strip()
    if level in LEVEL_KEYS:
        qs = qs.filter(student__level=level)
    if programme in PROGRAMME_KEYS:
        qs = qs.filter(student__programme=programme)
    if method:
        qs = qs.filter(method=method)
    if q:
        qs = qs.filter(
            Q(reference__icontains=q)
            | Q(receipt_no__icontains=q)
            | Q(phone_number__icontains=q)
            | Q(student__student_id__icontains=q)
            | Q(student__user__first_name__icontains=q)
            | Q(student__user__last_name__icontains=q)
        )
    return qs


def _xlsx(rows, headers, filename, title="Sheet1"):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF", name="Arial")
        cell.fill = PatternFill("solid", start_color="0B2E59")
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append(row)
    for i, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[i - 1])) for r in rows] + [10]) + 4
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(width, 45)
    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


@staff_required
def admin_dashboard(request):
    payments = Payment.objects.all()
    total_collected = payments.aggregate(t=Sum("amount"))["t"] or Decimal("0")
    today = timezone.localdate()
    today_total = payments.filter(created_at__date=today).aggregate(t=Sum("amount"))["t"] or Decimal("0")

    by_level = []
    for value, label in LEVELS:
        agg = payments.filter(student__level=value).aggregate(t=Sum("amount"), c=Count("id"))
        expected = Due.objects.filter(student__level=value).aggregate(t=Sum("amount"))["t"] or Decimal("0")
        collected = agg["t"] or Decimal("0")
        by_level.append(
            {
                "label": label,
                "value": value,
                "students": Student.objects.filter(level=value).count(),
                "payments": agg["c"],
                "collected": collected,
                "expected": expected,
                "outstanding": expected - collected,
                "percent": int(collected / expected * 100) if expected else 0,
            }
        )

    by_programme = []
    for value, label in PROGRAMMES:
        agg = payments.filter(student__programme=value).aggregate(t=Sum("amount"), c=Count("id"))
        by_programme.append(
            {
                "label": label,
                "students": Student.objects.filter(programme=value).count(),
                "payments": agg["c"],
                "collected": agg["t"] or Decimal("0"),
            }
        )

    expected_all = Due.objects.aggregate(t=Sum("amount"))["t"] or Decimal("0")
    return render(
        request,
        "admin_portal/dashboard.html",
        {
            "active": "dashboard",
            "total_collected": total_collected,
            "today_total": today_total,
            "expected_all": expected_all,
            "outstanding_all": expected_all - total_collected,
            "student_count": Student.objects.count(),
            "payment_count": payments.count(),
            "by_level": by_level,
            "by_programme": by_programme,
            "recent": payments.select_related("student", "student__user")[:8],
            "max_level_total": max([l["collected"] for l in by_level] or [Decimal("0")]) or Decimal("1"),
        },
    )


@staff_required
def admin_students(request):
    students = _filtered_students(request)
    return render(
        request,
        "admin_portal/students.html",
        {
            "active": "students",
            "students": students,
            "levels": LEVELS,
            "programmes": PROGRAMMES,
            "f_level": request.GET.get("level", ""),
            "f_programme": request.GET.get("programme", ""),
            "q": request.GET.get("q", ""),
            "default_password": DEFAULT_PASSWORD,
        },
    )


@staff_required
def admin_students_export(request):
    ids = request.POST.getlist("selected") if request.method == "POST" else []
    students = _filtered_students(request)
    if ids:
        students = Student.objects.select_related("user").filter(pk__in=ids)
    rows = [
        [
            s.full_name,
            s.student_id,
            s.level,
            s.programme_label,
            s.user.email,
            s.phone,
            float(s.total_dues),
            float(s.total_paid),
            float(s.outstanding),
        ]
        for s in students
    ]
    return _xlsx(
        rows,
        ["Name", "Index Number", "Level", "Programme", "Email", "Phone",
         "Total Dues (GHS)", "Paid (GHS)", "Outstanding (GHS)"],
        f"uenr-students-{timezone.localdate()}.xlsx",
        "Students",
    )


@staff_required
def admin_students_template(request):
    rows = [
        ["Kwame Mensah", "202312345", "100", "Computer Engineering", "kwame@uenr.edu.gh", "0244000000"],
        ["Ama Boateng", "202398765", "200", "Electrical and Electronics Engineering", "ama@uenr.edu.gh", "0201111111"],
    ]
    return _xlsx(
        rows,
        ["Name", "Index Number", "Level", "Programme", "Email", "Phone"],
        "uenr-student-upload-template.xlsx",
        "Template",
    )


def _read_upload(uploaded):
    """Yield dict rows from an .xlsx or .csv upload."""
    name = uploaded.name.lower()
    if name.endswith(".csv"):
        import csv
        import io

        text = uploaded.read().decode("utf-8-sig", errors="ignore")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
    else:
        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(uploaded.read()), data_only=True)
        rows = [[c if c is not None else "" for c in r] for r in wb.active.iter_rows(values_only=True)]
    if not rows:
        return []
    header = [str(c).strip().lower() for c in rows[0]]

    def idx(*names):
        for n in names:
            for i, h in enumerate(header):
                if n in h:
                    return i
        return None

    cols = {
        "name": idx("name"),
        "student_id": idx("index", "student id", "id"),
        "level": idx("level"),
        "programme": idx("programme", "program", "course"),
        "email": idx("email"),
        "phone": idx("phone", "mobile"),
    }
    out = []
    for row in rows[1:]:
        if not any(str(c).strip() for c in row):
            continue
        item = {}
        for key, i in cols.items():
            item[key] = str(row[i]).strip() if i is not None and i < len(row) and row[i] != "" else ""
        out.append(item)
    return out


@staff_required
def admin_students_upload(request):
    if request.method != "POST" or not request.FILES.get("file"):
        messages.error(request, "Choose an Excel (.xlsx) or CSV file to upload.")
        return redirect("admin_students")
    created = updated = skipped = 0
    try:
        rows = _read_upload(request.FILES["file"])
    except Exception as exc:  # noqa: BLE001
        messages.error(request, f"Could not read that file: {exc}")
        return redirect("admin_students")

    for row in rows:
        student_id = row.get("student_id", "")
        if not student_id:
            skipped += 1
            continue
        level = _level_from_text(row.get("level")) or "100"
        programme = _programme_from_text(row.get("programme"))
        first, _, last = (row.get("name") or "").partition(" ")
        student = Student.objects.filter(student_id=student_id).first()
        if student:
            student.level = level
            student.programme = programme or student.programme
            student.phone = row.get("phone") or student.phone
            student.save()
            if row.get("name"):
                student.user.first_name, student.user.last_name = first, last
            if row.get("email"):
                student.user.email = row["email"]
            student.user.save()
            updated += 1
            continue
        user = User.objects.filter(username=student_id).first()
        if user is None:
            user = User.objects.create_user(
                username=student_id, email=row.get("email") or "",
                password=DEFAULT_PASSWORD, first_name=first, last_name=last,
            )
        student = Student.objects.create(
            user=user, student_id=student_id, level=level,
            programme=programme, phone=row.get("phone") or "",
        )
        for template in DueTemplate.objects.filter(level=level):
            if template.programme and programme and template.programme != programme:
                continue
            Due.objects.get_or_create(
                student=student, template=template,
                defaults={"description": template.description, "amount": template.amount},
            )
        created += 1

    messages.success(
        request,
        f"Upload complete — {created} added, {updated} updated, {skipped} skipped. "
        f"New accounts use the default password {DEFAULT_PASSWORD}.",
    )
    return redirect("admin_students")


@staff_required
def admin_student_delete(request, pk):
    student = get_object_or_404(Student, pk=pk)
    if request.method == "POST":
        student.user.delete()
        messages.success(request, "Student removed.")
    return redirect("admin_students")


@staff_required
def admin_dues(request):
    if request.method == "POST":
        description = request.POST.get("description", "").strip()
        amount = request.POST.get("amount", "0")
        level = request.POST.get("level", "")
        programme = request.POST.get("programme", "")
        try:
            amount = Decimal(amount)
        except Exception:  # noqa: BLE001
            amount = Decimal("0")
        if not description or amount <= 0 or level not in LEVEL_KEYS:
            messages.error(request, "Enter a description, an amount above zero and a level.")
        else:
            template = DueTemplate.objects.create(
                description=description, amount=amount, level=level,
                programme=programme if programme in PROGRAMME_KEYS else "",
            )
            count = template.publish()
            messages.success(
                request,
                f"“{description}” published to {count} student(s) in Level {level}.",
            )
        return redirect("admin_dues")

    templates = DueTemplate.objects.all()
    for t in templates:
        t.assigned = t.dues.count()
        t.collected = t.dues.aggregate(x=Sum("paid"))["x"] or Decimal("0")
    return render(
        request,
        "admin_portal/dues.html",
        {"active": "dues", "templates": templates, "levels": LEVELS, "programmes": PROGRAMMES},
    )


@staff_required
def admin_due_republish(request, pk):
    template = get_object_or_404(DueTemplate, pk=pk)
    count = template.publish()
    messages.success(request, f"Added to {count} more student(s).")
    return redirect("admin_dues")


@staff_required
def admin_due_delete(request, pk):
    template = get_object_or_404(DueTemplate, pk=pk)
    if request.method == "POST":
        template.delete()
        messages.success(request, "Due removed from the platform.")
    return redirect("admin_dues")


@staff_required
def admin_payments(request):
    payments = _filtered_payments(request)
    total = payments.aggregate(t=Sum("amount"))["t"] or Decimal("0")
    from .models import PAYMENT_METHODS

    return render(
        request,
        "admin_portal/payments.html",
        {
            "active": "payments",
            "payments": payments,
            "total": total,
            "levels": LEVELS,
            "programmes": PROGRAMMES,
            "methods": PAYMENT_METHODS,
            "f_level": request.GET.get("level", ""),
            "f_programme": request.GET.get("programme", ""),
            "f_method": request.GET.get("method", ""),
            "q": request.GET.get("q", ""),
        },
    )


@staff_required
def admin_payments_export(request):
    payments = _filtered_payments(request)
    rows = [
        [
            p.student.full_name,
            p.student.student_id,
            p.student.level,
            p.student.programme_label,
            p.reference,
            p.receipt_no,
            p.phone_number,
            p.method_label,
            float(p.amount),
            timezone.localtime(p.created_at).strftime("%d %b %Y"),
            timezone.localtime(p.created_at).strftime("%I:%M %p"),
        ]
        for p in payments
    ]
    return _xlsx(
        rows,
        ["Name", "Index Number", "Level", "Programme", "Transaction Ref", "Receipt No",
         "Phone Number", "Method", "Amount (GHS)", "Date", "Time"],
        f"uenr-payments-{timezone.localdate()}.xlsx",
        "Payments",
    )


@staff_required
def admin_receipt(request, reference):
    from .views import _render_receipt

    payment = get_object_or_404(Payment, reference=reference)
    return _render_receipt(request, payment)
